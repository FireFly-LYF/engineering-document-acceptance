"""Create fictional archive packs with tables, true raster pages and independent truth."""
from __future__ import annotations
import argparse
import json
import sys
from pathlib import Path
import pymupdf
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.common import write_json

FONT = Path('C:/Windows/Fonts/simhei.ttf')
PROJECT = {
    'project_name':'澄川实验学校教学楼改扩建工程（模拟）','project_code':'CC-SCHOOL-2025-016',
    'owner':'澄川教育建设中心（虚构）','contractor':'澄川建筑有限公司（虚构）',
    'supervisor':'清源监理有限公司（虚构）','designer':'清源设计院（虚构）','surveyor':'澄川勘察院（虚构）',
    'address':'虚构市学苑路66号','contract_no':'SG-2025-016','contract_amount':'86500000',
    'start_date':'2025-02-18','finish_date':'2026-08-20'}
LABELS = {'project_name':'工程名称','project_code':'项目编号','owner':'建设单位','contractor':'施工单位',
    'supervisor':'监理单位','designer':'设计单位','surveyor':'勘察单位','address':'建设地点',
    'contract_no':'合同编号','contract_amount':'合同金额','start_date':'实际开工日期','finish_date':'实际竣工日期'}


def pdf(path,title,fields,*,rows=(),extra_pages=(),scan_pages=()):
    path = Path(path)
    path.parent.mkdir(parents=True,exist_ok=True)
    doc = pymupdf.open()
    payloads = [(title,fields,list(rows)),*extra_pages]
    for n,(heading,data,extra) in enumerate(payloads,1):
        page = doc.new_page(width=595,height=842)
        page.insert_font(fontname='cn',fontfile=str(FONT))
        page.insert_text((48,50),heading,fontname='cn',fontsize=16)
        page.insert_text((48,76),'虚构材料 仅用于系统测试 不作为真实工程文件',fontname='cn',fontsize=9,color=(0.55,0.55,0.55))
        y = 108
        for key,value in data.items():
            if y>710:
                raise ValueError('测试字段溢出页面')
            page.draw_rect(pymupdf.Rect(44,y-16,552,y+7),color=(0.45,0.45,0.45),width=.4)
            page.insert_text((50,y),f'{LABELS.get(key,key)}：{value}',fontname='cn',fontsize=11)
            y+=26
        for line in extra:
            page.insert_text((50,y),line,fontname='cn',fontsize=10)
            y+=25
        page.insert_text((50,794),f'测试页 {n} / {len(payloads)}',fontname='cn',fontsize=9)
        if n in scan_pages:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.6,1.6),alpha=False)
            replacement = pymupdf.open()
            replacement.new_page(width=595,height=842).insert_image(pymupdf.Rect(0,0,595,842),stream=pixmap.tobytes('png'))
            doc.delete_page(n-1)
            doc.insert_pdf(replacement,start_at=n-1)
            replacement.close()
    doc.subset_fonts()
    doc.save(path)
    doc.close()


def base(case,project=None,*,fire='special',use='school'):
    project = dict(project or PROJECT)
    manifest = {'project_id':case.name,**project,'project_type':'房屋建筑工程',
        'audit_context':{'audit_stage':'post_acceptance','regulatory_scope':'general_civil','fire_regulation_applicable':True,'use_type':use,'supervision_entrusted':True,
            'fire_category':fire,'fire_category_source':'测试人工确认条件，不由模型推断',
            'fire_spotcheck_selected':False,'coverage_plan_confirmed':False},
        'data_note':'全部项目单位和文件均为虚构模拟；真实像素扫描页无文本层'}
    case.mkdir(parents=True,exist_ok=True)
    files = case/'files'
    files.mkdir(exist_ok=True)
    common = {k:project[k] for k in ('project_name','project_code','owner','address')}
    return manifest,files,common


def required(files,common,project=PROJECT,*,formal=True,supervision_code=None,amount=None,fire=True):
    pdf(files/'01-准备'/'合同.pdf','建设工程施工合同',{**common,'文件编号':'HT-016','contract_no':project['contract_no'],
        'contractor':project['contractor']},rows=['合同价款条款见第三页'],extra_pages=[
        ('合同履约与施工范围',{'单位工程名称':'教学楼'},['履约范围为主体及安装工程']),
        ('合同价格与计价',{**common,'contract_no':project['contract_no'],'合同金额（万元）':'8650'},[])])
    title = '工程竣工验收报告' if formal else '工程竣工预验收报告'
    pdf(files/'05-验收'/'报告.pdf',title,{**common,'文件编号':'YS-016','验收日期':'2026-08-28','验收结论':'文载结论为合格（模拟）'})
    pdf(files/'05-验收'/'施工报告.pdf','施工单位工程竣工报告',{**common,'文件编号':'SG-JG-016',
        'contract_no':project['contract_no'],'contractor':project['contractor'],'contract_amount':amount or project['contract_amount'],
        'start_date':project['start_date'],'finish_date':project['finish_date']})
    pdf(files/'02-监理'/'质量报告.pdf','监理工程质量评估报告',{**common,
        'project_code':supervision_code or project['project_code'],'文件编号':'JL-016','supervisor':project['supervisor'],
        '评估结论':'同意进入竣工验收程序（模拟）'},rows=['附件引用 项目编号：'+project['project_code']])
    pdf(files/'05-验收'/'勘察检查.pdf','勘察单位工程质量检查报告',{**common,'文件编号':'KC-016','surveyor':project['surveyor'],'检查结论':'文载合格（模拟）'})
    pdf(files/'05-验收'/'设计检查.pdf','设计单位工程质量检查报告',{**common,'文件编号':'SJ-016','designer':project['designer'],'检查结论':'文载合格（模拟）'})
    pdf(files/'05-验收'/'保修.pdf','工程质量保修书',{**common,'文件编号':'BX-016','contractor':project['contractor']})
    if fire:
        pdf(files/'07-专项'/'行政文书.pdf','特殊建设工程消防验收意见书',{**common,'消防验收意见书编号':'XF-016','验收结论':'合格（模拟）'})


def build(root):
    root = Path(root)
    created = []
    case = root/'school-issues'
    manifest,files,common = base(case)
    required(files,common,formal=False,supervision_code='CC-SCHOOL-2025-017',amount='92000000',fire=False)
    pdf(files/'07-专项'/'规划文件.pdf','建设工程规划核实合格证',{**common,'文件编号':'GH-016','建筑面积':'12600平方米'})
    pdf(files/'07-专项'/'申请材料.pdf','建设工程消防验收备案凭证',{**common,'消防备案凭证编号':'BA-016'})
    other = {'project_name':'澄川河堤整治工程（模拟）','project_code':'CC-RIVER-2024-008','owner':'澄川水务建设中心（虚构）','address':'虚构市滨水路8号'}
    pdf(files/'01-准备'/'补充合同.pdf','建设工程施工合同',{**other,'文件编号':'HT-008','合同编号':'SG-2024-008','合同金额':'120000000元'})
    pdf(files/'05-验收'/'补充报告.pdf','工程竣工验收报告',{**other,'文件编号':'YS-008'})
    record = files/'03-施工'/'记录.pdf'
    pdf(record,'检验批质量验收记录',{**common,'文件编号':'JYP-001','单位工程名称':'教学楼','专业':'结构','工程部位':'一层1-5轴','验收日期':'2025-04-16'})
    import shutil
    shutil.copy2(record,files/'03-施工'/'记录备份.pdf')
    write_json(case/'manifest.json',manifest)
    write_json(case/'ground_truth.json',{'missing_rules':['AC-001','FI-001'],'misfiled':2,'field_conflicts':2,'duplicate_groups':1,'third_page_contract_amount':'86500000'})
    created.append(case)

    case = root/'school-complete'
    manifest,files,common = base(case)
    required(files,common)
    write_json(case/'manifest.json',manifest)
    write_json(case/'ground_truth.json',{'missing_rules':[],'field_conflicts':0,'misfiled':0})
    created.append(case)

    case = root/'residential-coverage'
    project = {**PROJECT,'project_name':'澄川景苑一期住宅工程（模拟）','project_code':'CC-HOME-2025-021'}
    manifest,files,common = base(case,project,fire='other',use='residential')
    required(files,common,project,fire=False)
    pdf(files/'07-专项'/'备案文书.pdf','建设工程消防验收备案凭证',{**common,'消防备案凭证编号':'BA-021'})
    pdf(files/'05-验收'/'分户验收.pdf','住宅工程质量分户验收表',{**common,'单位工程名称':'1号楼','工程部位':'101室'})
    pdf(files/'03-施工'/'检验记录.pdf','检验批质量验收记录',{**common,'文件编号':'JYP-021','单位工程名称':'1号楼','专业':'结构','工程部位':'一层'})
    manifest['audit_context']['coverage_plan_confirmed'] = True
    manifest['coverage_plan'] = [{'tag':'INSPECTION_BATCH','unit_name':'1号楼','discipline':'结构','location':'一层'},
        {'tag':'INSPECTION_BATCH','unit_name':'2号楼','discipline':'结构','location':'一层'}]
    write_json(case/'manifest.json',manifest)
    write_json(case/'ground_truth.json',{'coverage_gaps':1,'missing_rules':[]})
    created.append(case)

    case = root/'versions-and-office'
    manifest,files,common = base(case)
    required(files,common)
    from docx import Document
    document = Document()
    document.add_heading('设计变更通知单',0)
    for key,value in common.items():
        document.add_paragraph(LABELS.get(key,key)+'：'+value)
    document.add_paragraph('设计变更编号：BG-006')
    document.add_paragraph('变更后合同金额：9200万元')
    document.add_paragraph('结算金额：9300万元')
    target = files/'03-施工'/'变更.docx'
    target.parent.mkdir(exist_ok=True)
    document.save(target)
    for revision in ('A','B'):
        pdf(files/'06-竣工图'/f'结构图-{revision}.pdf','结构竣工图说明',{**common,'文件编号':'JG-01','图号':'S-01','专业':'结构','版本号':revision})
    import openpyxl
    book = openpyxl.Workbook()
    book.active.title = '封面'
    book.active.append(['移交资料说明'])
    sheet = book.create_sheet('竣工目录')
    sheet.append(['竣工资料目录'])
    for key,value in common.items():
        sheet.append([LABELS.get(key,key)+'：'+value])
    sheet.append(['专业：结构','图号：S-01'])
    (files/'08-移交').mkdir(exist_ok=True)
    book.save(files/'08-移交'/'移交目录.xlsx')
    write_json(case/'manifest.json',manifest)
    write_json(case/'ground_truth.json',{'contract_amount_conflicts':0,'version_conflicts':1,'second_sheet_read':True})
    created.append(case)

    case = root/'scans-and-unreadable'
    manifest,files,common = base(case)
    required(files,common)
    manifest['audit_context'].pop('fire_category')
    pdf(files/'04-检测'/'扫描报告.pdf','钢筋复试报告',{**common,'报告编号':'JC-001','材料名称':'热轧带肋钢筋','规格型号':'HRB400E','样品编号':'GJ-001','检测日期':'2025-03-18'},scan_pages=(1,))
    pdf(files/'03-施工'/'混合记录.pdf','隐蔽工程验收记录',{**common,'文件编号':'YB-001'},extra_pages=[
        ('隐蔽工程验收记录',{**common,'工程部位':'二层梁柱节点','验收日期':'2025-05-20'},[])],scan_pages=(2,))
    (files/'04-检测'/'补充检测.pdf').write_bytes(b'%PDF-1.7\nBROKEN TEST FILE')
    pdf(files/'05-验收'/'合订报告.pdf','设计单位工程质量检查报告',{**common,'文件编号':'SJ-100'},extra_pages=[
        ('勘察单位工程质量检查报告',{**common,'文件编号':'KC-100'},[])])
    write_json(case/'manifest.json',manifest)
    write_json(case/'ground_truth.json',{'offline_unreadable_min':3,'fire_rules_unknown':True,'scanned_pages':[1,2],'binder_documents':2})
    created.append(case)
    return created


if __name__=='__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--root',default='cases/realistic')
    parser.add_argument('--missing-only',action='store_true',help='为仓库已有测试元数据补齐缺失的合成文件，不覆盖任何已有文件')
    args = parser.parse_args()
    target = Path(args.root)
    if args.missing_only:
        import shutil
        import uuid
        project_root = Path(__file__).resolve().parents[1]
        staging = project_root/'tmp'/('fixture-build-'+uuid.uuid4().hex[:12])
        generated = build(staging)
        # Validate metadata before any copying. Changed/user-authored cases stay untouched.
        for source_case in generated:
            for filename in ('manifest.json','ground_truth.json'):
                destination = target/source_case.name/filename
                if destination.exists() and json.loads(destination.read_text(encoding='utf-8-sig'))!=json.loads((source_case/filename).read_text(encoding='utf-8')):
                    raise SystemExit('已有测试元数据与生成器不同，请使用新目录，避免覆盖材料')
        copied = 0
        for source_case in generated:
            for source in source_case.rglob('*'):
                if source.is_file():
                    destination = target/source_case.name/source.relative_to(source_case)
                    if not destination.exists():
                        destination.parent.mkdir(parents=True,exist_ok=True)
                        shutil.copy2(source,destination)
                        copied+=1
        print('Missing synthetic files restored:',copied)
        raise SystemExit(0)
    if target.exists() and any(target.iterdir()):
        raise SystemExit('目标目录非空；请选择新目录，避免覆盖测试或用户材料')
    cases = build(target)
    print('Generated fictional cases:',len(cases))
